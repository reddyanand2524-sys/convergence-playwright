"""

Builds an openpyxl workbook entirely in-memory (io.BytesIO, never written to
disk) with sheets: Summary, Tests, Steps, Failures, Screenshots (metadata
only — no embedded images).

Security: any user-controlled string written to a cell is passed through
_sanitize_cell_value() first, which neutralizes formula-injection payloads
(values starting with =, +, -, @) by prefixing a leading apostrophe so
Excel/LibreOffice treat the cell as literal text, never a formula.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.services.report_data_service import ReportData

_DANGEROUS_PREFIXES = ("=", "+", "-", "@")

_FILL_PASS = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_FILL_FAIL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_FILL_SKIP = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
_HEADER_FILL = PatternFill(start_color="404040", end_color="404040", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")

_STATUS_FILLS = {
    "PASSED": _FILL_PASS,
    "PASS": _FILL_PASS,
    "FAILED": _FILL_FAIL,
    "FAIL": _FILL_FAIL,
    "ERROR": _FILL_FAIL,
    "SKIPPED": _FILL_SKIP,
}


def _sanitize_cell_value(value):
    """Neutralize formula-injection payloads for any text destined for a
    spreadsheet cell. Non-string values pass through unchanged."""
    if isinstance(value, str) and value.startswith(_DANGEROUS_PREFIXES):
        return "'" + value
    return value


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width


def _failed_step_text(analysis) -> str:
    step = analysis.failed_step
    if step is None:
        return ""
    number = step.step_number if step.step_number is not None else step.sequence_index
    return f"#{number} {step.action}" + (f" {step.target}" if step.target else "")


def render_excel_report(report: ReportData) -> bytes:
    wb = Workbook()

    # --- Summary sheet ---
    ws_summary = wb.active
    ws_summary.title = "Summary"
    _write_header(ws_summary, ["Field", "Value"])
    rows = [
        ("Project", _sanitize_cell_value(report.project.name)),
        ("Project ID", report.project.id),
        ("Run ID", report.run.id),
        ("Browser", report.run.browser),
        ("Status", report.run.status),
        ("Created At", str(report.run.created_at) if report.run.created_at else ""),
        ("Started At", str(report.run.started_at) if report.run.started_at else ""),
        ("Completed At", str(report.run.completed_at) if report.run.completed_at else ""),
        ("Run Duration (s)", report.run.duration_seconds),
        ("Cancel Requested", report.run.cancel_requested),
        ("Total Tests", report.summary.total_tests),
        ("Passed Tests", report.summary.passed_count),
        ("Failed Tests", report.summary.failed_count),
        ("Skipped Tests", report.summary.skipped_count),
        ("Total Steps", report.summary.total_steps),
        ("Passed Steps", report.summary.passed_steps),
        ("Failed Steps", report.summary.failed_steps),
        ("Screenshot Count", report.summary.screenshot_count),
        ("Failure Count", report.summary.failure_count),
        ("Test Duration (s)", report.summary.duration_seconds),
        ("Exit Code", report.summary.exit_code),
        ("Notice", _sanitize_cell_value(report.notice) if report.notice else ""),
    ]
    for field_name, value in rows:
        ws_summary.append([field_name, value])
    _autosize(ws_summary, [24, 40])

    # --- Tests sheet ---
    ws_tests = wb.create_sheet("Tests")
    _write_header(
        ws_tests,
        ["ID", "Node ID", "File", "Class", "Method", "Status", "Duration (s)"],
    )
    for test in report.tests:
        row_idx = ws_tests.max_row + 1
        ws_tests.append(
            [
                test.id,
                _sanitize_cell_value(test.nodeid),
                _sanitize_cell_value(test.file_name),
                _sanitize_cell_value(test.class_name) if test.class_name else "",
                _sanitize_cell_value(test.method_name),
                test.status,
                test.duration_seconds,
            ]
        )
        fill = _STATUS_FILLS.get(test.status)
        if fill:
            for cell in ws_tests[row_idx]:
                cell.fill = fill
    _autosize(ws_tests, [8, 45, 25, 20, 25, 12, 14])

    # --- Steps sheet ---
    ws_steps = wb.create_sheet("Steps")
    _write_header(
        ws_steps,
        [
            "Step ID",
            "Test ID",
            "Seq",
            "Action",
            "Target",
            "Value",
            "Status",
            "Duration (s)",
            "Error",
        ],
    )
    for test in report.tests:
        for step in test.steps:
            row_idx = ws_steps.max_row + 1
            ws_steps.append(
                [
                    step.id,
                    test.id,
                    step.sequence_index,
                    _sanitize_cell_value(step.action),
                    _sanitize_cell_value(step.target) if step.target else "",
                    _sanitize_cell_value(step.value) if step.value else "",
                    step.status,
                    step.duration_seconds,
                    _sanitize_cell_value(step.error) if step.error else "",
                ]
            )
            fill = _STATUS_FILLS.get(step.status)
            if fill:
                for cell in ws_steps[row_idx]:
                    cell.fill = fill
            ws_steps.cell(row=row_idx, column=9).alignment = _WRAP
    _autosize(ws_steps, [8, 8, 6, 20, 30, 20, 10, 12, 50])

    # --- Failures sheet ---
    ws_failures = wb.create_sheet("Failures")
    _write_header(
        ws_failures,
        [
            "Test ID",
            "Node ID",
            "Exception Type",
            "Exception Message",
            "Source File",
            "Source Line Number",
            "Source Line",
            "Failure Category",
            "Confidence",
            "Summary",
            "Likely Cause (hypothesis)",
            "Evidence",
            "Recommended Actions",
            "Failed Step",
        ],
    )
    for test in report.tests:
        if not test.failure and not test.analysis:
            continue
        f = test.failure
        a = test.analysis
        row_idx = ws_failures.max_row + 1
        ws_failures.append(
            [
                test.id,
                _sanitize_cell_value(test.nodeid),
                _sanitize_cell_value(f.exception_type if f else a.exception_type) or "",
                _sanitize_cell_value(a.exception_message if a else f.exception_message) or "",
                _sanitize_cell_value(f.source_file if f else a.file_name) or "",
                f.source_line_number if f else a.line_number,
                _sanitize_cell_value(a.source_line if a else f.source_line) or "",
                a.failure_category if a else "",
                a.confidence if a else "",
                _sanitize_cell_value(a.summary) if a else "",
                _sanitize_cell_value(a.likely_cause) if a else "",
                _sanitize_cell_value(chr(10).join(f"{e.label}: {e.value}" for e in a.evidence)) if a else "",
                _sanitize_cell_value(chr(10).join(f"{i}. {x}" for i, x in enumerate(a.recommended_actions, 1))) if a else "",
                _sanitize_cell_value(_failed_step_text(a)) if a else "",
            ]
        )
        for col in (4, 7, 10, 11, 12, 13, 14):
            ws_failures.cell(row=row_idx, column=col).alignment = _WRAP
    _autosize(ws_failures, [8, 45, 20, 50, 25, 10, 40, 24, 12, 50, 60, 60, 60, 40])

    # --- Screenshots sheet (metadata only, no embedded images) ---
    ws_shots = wb.create_sheet("Screenshots")
    _write_header(
        ws_shots,
        ["Screenshot ID", "Test ID", "Step ID", "File Name", "Mime Type", "Capture Reason", "Size (bytes)"],
    )
    for test in report.tests:
        for shot in test.screenshots:
            ws_shots.append(
                [
                    shot.id,
                    test.id,
                    "",
                    _sanitize_cell_value(shot.file_name),
                    shot.mime_type,
                    shot.capture_reason,
                    shot.size_bytes,
                ]
            )
        for step in test.steps:
            for shot in step.screenshots:
                ws_shots.append(
                    [
                        shot.id,
                        test.id,
                        step.id,
                        _sanitize_cell_value(shot.file_name),
                        shot.mime_type,
                        shot.capture_reason,
                        shot.size_bytes,
                    ]
                )
    _autosize(ws_shots, [12, 8, 8, 30, 14, 16, 12])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
