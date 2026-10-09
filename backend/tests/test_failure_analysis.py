"""Failure analysis: classification, evidence, redaction, API and reports.

Pure-logic tests use no database. API/report tests use the shared
SAVEPOINT-rollback db_session fixture, so nothing is committed.
"""
import io
import json
from datetime import datetime, timedelta
from xml.sax.saxutils import escape

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.security import create_access_token, hash_password
from app.models.failure_diagnostic import FailureDiagnostic
from app.models.screenshot import Screenshot
from app.models.step_result import StepResult
from app.models.test_result import TestResult
from app.models.user import User
from app.services import failure_analysis_rules as rules
from app.services import failure_analysis_service as fas
from app.services import failure_sanitizer as san
from app.services.failure_analysis_service import (
    FailureInput, ScreenshotInput, StepInput, analyze_failure, classify, safe_analyze,
)
from app.services.report_data_service import build_report_data
from app.services.report_excel_service import render_excel_report
from app.services.report_html_service import render_html_report
from app.services.report_pdf_service import render_pdf_report
from app.services.result_persistence_service import persist_run_results
from tests.conftest import make_project_file, make_run

NOW = datetime(2026, 10, 9, 12, 0, 0)


def make_input(exc_type="selenium.common.exceptions.NoSuchElementException", message="boom", **kw):
    base = dict(
        test_id=1, nodeid="test_login.py::test_x", status="FAILED", has_failure_record=True,
        file_name="test_login.py", method_name="test_x", duration_seconds=3.2, completed_at=NOW,
        browser="Chrome", exception_type=exc_type, exception_message=message,
        source_file="test_login.py", source_line_number=10, source_line='driver.find_element(By.ID, "x").click()',
    )
    base.update(kw)
    return FailureInput(**base)


def step(i, action="click", target="id=a", status="PASS", shots=(), duration=0.2, error=None, at=None):
    return StepInput(
        sequence_index=i, step_number=i, action=action, target=target, status=status,
        duration_seconds=duration, error=error, occurred_at=at, screenshots=tuple(shots),
    )


# ---------------------------------------------------------------------------
# classification: every supported category
# ---------------------------------------------------------------------------

CASES = [
    ("selenium.common.exceptions.NoSuchElementException", 'no such element: Unable to locate element: {"method":"css selector","selector":"#nope"}', rules.ELEMENT_NOT_FOUND, "high"),
    ("selenium.common.exceptions.TimeoutException", "Message: ", rules.TIMEOUT, "high"),
    ("selenium.common.exceptions.ElementNotInteractableException", "element not interactable", rules.ELEMENT_NOT_INTERACTABLE, "high"),
    ("selenium.common.exceptions.ElementClickInterceptedException", "element click intercepted: Element <a> is not clickable at point (10, 10). Other element would receive the click: <div class=\"overlay\"></div>", rules.CLICK_INTERCEPTED, "high"),
    ("selenium.common.exceptions.StaleElementReferenceException", "stale element not found", rules.STALE_ELEMENT, "high"),
    ("AssertionError", "assert 'Welcome' == 'Login'", rules.ASSERTION_FAILURE, "high"),
    ("assert False", "", rules.ASSERTION_FAILURE, "high"),          # pytest's rendering of a plain assert
    ("assert 1 == 2", None, rules.ASSERTION_FAILURE, "high"),
    (None, "assert 'a' == 'b'", rules.ASSERTION_FAILURE, "high"),   # type missing, message is the assert text
    ("Failed", "explicit pytest.fail", rules.ASSERTION_FAILURE, "medium"),
    ("selenium.common.exceptions.InvalidArgumentException", "invalid argument: invalid 'url'", rules.NAVIGATION_FAILURE, "high"),
    ("selenium.common.exceptions.WebDriverException", "unknown error: net::ERR_NAME_NOT_RESOLVED", rules.NETWORK_FAILURE, "medium"),
    ("selenium.common.exceptions.WebDriverException", "unknown error: net::ERR_ABORTED", rules.NAVIGATION_FAILURE, "medium"),
    ("selenium.common.exceptions.WebDriverException", "Message: Internal Server Error returned by the page", rules.APPLICATION_ERROR, "medium"),
    ("requests.exceptions.HTTPError", "500 Server Error: for url: http://x", rules.API_ERROR, "high"),
    ("selenium.common.exceptions.SessionNotCreatedException", "session not created: Chrome instance exited", rules.BROWSER_CONFIGURATION_ERROR, "high"),
    ("selenium.common.exceptions.WebDriverException", "disconnected: not connected to DevTools", rules.WEBDRIVER_ERROR, "medium"),
    ("selenium.common.exceptions.WebDriverException", "something odd happened", rules.WEBDRIVER_ERROR, "medium"),
    ("KeyError", "'username'", rules.TEST_DATA_ERROR, "low"),
    ("json.decoder.JSONDecodeError", "Expecting value", rules.TEST_DATA_ERROR, "low"),
    ("ModuleNotFoundError", "No module named 'pandas'", rules.ENVIRONMENT_ERROR, "low"),
    ("builtins.RuntimeError", "completely unrecognised", rules.UNKNOWN, "low"),
]


@pytest.mark.parametrize("exc_type,message,expected,confidence", CASES)
def test_each_category_is_classified(exc_type, message, expected, confidence):
    category, conf, _ = classify(exc_type, message, None)
    assert (category, conf) == (expected, confidence)


def test_every_declared_category_is_covered_by_cases_or_rules():
    covered = {c[2] for c in CASES}
    assert covered == set(rules.CATEGORIES)
    assert set(rules.TITLES) == set(rules.LIKELY_CAUSES) == set(rules.RECOMMENDED_ACTIONS) == set(rules.CATEGORIES)


def test_timeout_is_not_always_a_slow_page():
    # Same exception class, different context -> different category, lower confidence.
    plain = classify("selenium.common.exceptions.TimeoutException", "", step(1, "click"))
    nav = classify("selenium.common.exceptions.TimeoutException", "", step(1, "open", "https://x.test"))
    assert plain[0] == rules.TIMEOUT and nav[0] == rules.NAVIGATION_FAILURE
    assert nav[1] == "medium" and nav[2] == "exception_type+failed_step"


def test_invalid_argument_without_url_is_not_called_navigation():
    category, conf, _ = classify("selenium.common.exceptions.InvalidArgumentException", "invalid argument: bad keys", None)
    assert (category, conf) == (rules.WEBDRIVER_ERROR, "medium")


# ---------------------------------------------------------------------------
# honesty: facts vs hypotheses, no invented values, no sleep advice
# ---------------------------------------------------------------------------

def test_evidence_is_recorded_facts_and_cause_is_hedged():
    a = analyze_failure(make_input(
        message='no such element: Unable to locate element: {"method":"css selector","selector":"#nope"}',
        steps=(step(1, "open", "https://app.test/login"), step(2, "click", "id=go", status="FAIL", duration=22.5)),
    ))
    facts = {e.label: e.value for e in a.evidence}
    assert facts["Locator"] == "css selector: #nope"
    assert facts["Exception"].endswith("NoSuchElementException")
    assert "#2 click id=go (ran for 22.5s)" == facts["Failed step"]
    assert any(w in a.likely_cause.lower() for w in ("may", "possible"))
    assert a.category_basis == "exception_type" and a.confidence == "high"
    # recommendations are not mixed into evidence
    assert not set(a.recommended_actions) & {e.value for e in a.evidence}


@pytest.mark.parametrize("selector_text", [
    '[id="does-not-exist"]',      # real Selenium output: inner quotes are NOT escaped
    '[id=\\"does-not-exist\\"]',  # tolerated: escaped form
])
def test_locator_with_quotes_is_shown_in_full(selector_text):
    message = 'Message: no such element: Unable to locate element: {"method":"css selector","selector":"%s"}' % selector_text
    a = analyze_failure(make_input(message=message))
    assert next(e.value for e in a.evidence if e.label == "Locator") == 'css selector: [id="does-not-exist"]'
    assert 'does-not-exist' in a.summary


def test_assertion_values_are_only_shown_if_recorded():
    recorded = analyze_failure(make_input("AssertionError", "assert 'Welcome' == 'Login'"))
    empty = analyze_failure(make_input("AssertionError", None))
    assert next(e.value for e in recorded.evidence if e.label == "Message") == "assert 'Welcome' == 'Login'"
    assert "Message" not in {e.label for e in empty.evidence}
    assert "expected" not in empty.summary.lower() and "actual" not in empty.summary.lower()


@pytest.mark.parametrize("category", rules.CATEGORIES)
def test_recommendations_never_suggest_sleeping(category):
    for action in rules.RECOMMENDED_ACTIONS[category]:
        assert "time.sleep" not in action


def test_click_intercepted_evidence_names_the_blocking_element():
    a = analyze_failure(make_input(
        "selenium.common.exceptions.ElementClickInterceptedException",
        'Other element would receive the click: <div class="cookie-banner"></div>',
    ))
    assert any(e.label == "Element that received the click" and "cookie-banner" in e.value for e in a.evidence)


# ---------------------------------------------------------------------------
# incomplete / legacy / malformed data
# ---------------------------------------------------------------------------

def test_passing_test_has_no_analysis():
    assert analyze_failure(make_input(status="PASSED", has_failure_record=False)) is None


def test_legacy_failed_test_without_diagnostic_record_uses_step_error():
    inp = FailureInput(
        test_id=5, nodeid="t.py::t", status="FAILED", has_failure_record=False,
        steps=(step(1, "click", "id=x", status="FAIL", error="net::ERR_CONNECTION_REFUSED"),),
    )
    a = analyze_failure(inp)
    assert a.failure_category == rules.NETWORK_FAILURE and a.confidence == "medium"
    assert a.exception_type is None and a.traceback is None


def test_nothing_recorded_is_unknown_with_low_confidence():
    a = analyze_failure(FailureInput(test_id=1, nodeid="t.py::t", status="FAILED"))
    assert (a.failure_category, a.confidence) == (rules.UNKNOWN, "low")
    assert a.failed_step is None and a.timeline == () and a.screenshot_id is None
    assert {"Steps": "No steps were recorded for this test."}.items() <= {e.label: e.value for e in a.evidence}.items()


def test_missing_message_still_classifies_by_exception_type():
    a = analyze_failure(make_input("selenium.common.exceptions.StaleElementReferenceException", None))
    assert a.failure_category == rules.STALE_ELEMENT and a.exception_message is None


def test_malformed_values_do_not_crash():
    a = analyze_failure(make_input(source_line_number="not-a-number", exception_message=12345, source_line=b"x"))
    assert a.line_number is None and a.exception_message == "12345"


def test_error_status_is_explained_as_outside_the_test_body():
    a = analyze_failure(make_input("selenium.common.exceptions.SessionNotCreatedException", "session not created", status="ERROR"))
    facts = {e.label: e.value for e in a.evidence}
    assert a.failure_category == rules.BROWSER_CONFIGURATION_ERROR and "setup or teardown" in facts["Test status"]


def test_safe_analyze_never_raises(monkeypatch):
    monkeypatch.setattr(fas, "analyze_failure", lambda inp: (_ for _ in ()).throw(RuntimeError("bug")))
    assert safe_analyze(make_input()) is None


# ---------------------------------------------------------------------------
# file/line, timeline, page URL, screenshot choice
# ---------------------------------------------------------------------------

def test_location_is_taken_from_the_diagnostic():
    a = analyze_failure(make_input(source_file="test_login.py", source_line_number=42))
    assert (a.file_name, a.line_number) == ("test_login.py", 42)
    assert {"Location": "test_login.py:42"}.items() <= {e.label: e.value for e in a.evidence}.items()


def test_timeline_marks_failed_and_last_successful_steps():
    steps = tuple(step(i, "click", f"id=s{i}") for i in range(1, 9)) + (step(9, "click", "id=bad", status="FAIL"), step(10, "click", "id=never-ran", status="PASS"))
    a = analyze_failure(make_input(steps=steps))
    roles = [(t.sequence_index, t.role) for t in a.timeline]
    assert roles == [(5, "preceding"), (6, "preceding"), (7, "preceding"), (8, "last_successful"), (9, "failed")]
    assert a.failed_step.sequence_index == 9 and a.last_successful_step.sequence_index == 8


def test_failure_outside_a_recorded_step_has_no_failed_step():
    a = analyze_failure(make_input("AssertionError", "assert False", steps=(step(1), step(2))))
    assert a.failed_step is None and a.last_successful_step.sequence_index == 2
    # evidence states only the recorded fact, without guessing where it failed
    assert {"Failed step": "No recorded step is marked failed."}.items() <= {e.label: e.value for e in a.evidence}.items()


def test_page_url_is_last_navigation_before_failure_and_is_redacted():
    steps = (
        step(1, "open", "https://old.test/"),
        step(2, "open", "https://app.test/cb?token=abc123&page=2"),
        step(3, "click", "id=go", status="FAIL"),
        step(4, "open", "https://after.test/"),
    )
    a = analyze_failure(make_input(steps=steps))
    assert a.page_url == "https://app.test/cb?token=***&page=2"
    assert "not verified" in a.page_url_source


def test_screenshot_preference_order():
    failure_shot = ScreenshotInput(30, "TEST_FAILURE")
    step_shot = ScreenshotInput(20, "STEP_FAILURE")
    plain = ScreenshotInput(10, "ALL_MODE")
    s = (step(1, shots=[plain]), step(2, status="FAIL", shots=[step_shot]))
    assert analyze_failure(make_input(steps=s, test_screenshots=(failure_shot,))).screenshot_id == 30
    assert analyze_failure(make_input(steps=s)).screenshot_id == 20
    assert analyze_failure(make_input(steps=(step(1, shots=[plain]),))).screenshot_id == 10
    assert analyze_failure(make_input()).screenshot_id is None


# ---------------------------------------------------------------------------
# redaction
# ---------------------------------------------------------------------------

def test_secrets_and_paths_never_reach_the_analysis():
    inp = make_input(
        message="Message: bad login password=hunter2 Authorization: Bearer abcdefghijkl12345\nStacktrace:\n#0 0x1 <unknown>",
        source_line='driver.find_element(By.ID, "password").send_keys("hunter2")',
        raw_traceback="/tmp/run_9/test_login.py:10: in test_x\n    driver.find_element(By.ID, 'password').send_keys('hunter2')\nE   token=SECRETVALUE",
    )
    text = repr(analyze_failure(inp))
    for leaked in ("hunter2", "abcdefghijkl12345", "SECRETVALUE", "/tmp/run_9", "0x1 <unknown>"):
        assert leaked not in text
    assert "test_login.py:10" in text


@pytest.mark.parametrize("exc_type", [
    "assert d.title == 'x'",             # dot inside the expression
    "assert 'a.b' == 'c'",               # dotted string literal
    "assert driver.current_url.endswith('/home')",
])
def test_plain_assert_with_dots_is_an_assertion_failure(exc_type):
    assert classify(exc_type, None, None) == (rules.ASSERTION_FAILURE, "high", "pytest_assertion")


def test_literal_compared_with_a_secret_named_field_is_masked_but_ordinary_asserts_are_not():
    secret = "assert driver.find_element(By.ID, 'password').get_attribute('value') == 'hunter2-SECRET'"
    ordinary = "assert 'Welcome' == 'Login'"
    assert "hunter2-SECRET" not in san.clean_exception_type(secret)
    assert san.clean_exception_type(ordinary) == ordinary
    a = analyze_failure(make_input(secret, None))
    assert "hunter2-SECRET" not in repr(a) and a.failure_category == rules.ASSERTION_FAILURE


PYTEST_ASSERT_TRACE = (
    "test_x.py:5: in test_x\n"
    ">       assert d.find_element(By.ID, \"password\").get_attribute(\"value\") == \"CMP-CCC\"\n"
    "E   AssertionError: assert 'PAGE-AAA' == 'CMP-CCC'\n"
    "E   assert 'PAGE-AAA' == 'CMP-CCC'\n"
    "E     - CMP-CCC\n"
    "E     + PAGE-AAA"
)


def test_assertion_values_are_masked_only_when_the_failing_line_involves_a_secret_field():
    secret_line = 'assert d.find_element(By.ID, "password").get_attribute("value") == "CMP-CCC"'
    plain_line = 'assert d.title == "CMP-CCC"'

    masked = analyze_failure(make_input(
        "AssertionError", "assert 'PAGE-AAA' == 'CMP-CCC'", source_line=secret_line, raw_traceback=PYTEST_ASSERT_TRACE,
    ))
    plain = analyze_failure(make_input(
        "AssertionError", "assert 'PAGE-AAA' == 'CMP-CCC'", source_line=plain_line, raw_traceback=PYTEST_ASSERT_TRACE,
    ))

    assert "CMP-CCC" not in repr(masked) and "PAGE-AAA" not in repr(masked)
    assert masked.failure_category == rules.ASSERTION_FAILURE and "assert" in masked.exception_message
    # ordinary assertion output stays readable (the trace's ">" line is the only hint source)
    ordinary_trace = PYTEST_ASSERT_TRACE.replace('By.ID, "password"', 'By.ID, "title"')
    plain_no_hint = analyze_failure(make_input(
        "AssertionError", "assert 'PAGE-AAA' == 'CMP-CCC'", source_line=plain_line, raw_traceback=ordinary_trace,
    ))
    assert plain_no_hint.exception_message == "assert 'PAGE-AAA' == 'CMP-CCC'" and "CMP-CCC" in plain_no_hint.traceback
    assert plain.exception_message != plain_no_hint.exception_message  # same text, different trace context

    # source line empty (as pytest stores it for failed asserts): the ">" line is enough
    no_source = analyze_failure(make_input(
        "AssertionError", "assert 'PAGE-AAA' == 'CMP-CCC'", source_line=None, raw_traceback=PYTEST_ASSERT_TRACE,
    ))
    assert "CMP-CCC" not in repr(no_source) and "PAGE-AAA" not in repr(no_source)


def test_exception_type_text_is_sanitized_everywhere():
    secret_assert = "assert field.get_attribute('value') == 'x' and password = hunter2-SECRET"
    a = analyze_failure(make_input(secret_assert, None))
    text = repr(a)
    assert "hunter2-SECRET" not in text and a.failure_category == rules.ASSERTION_FAILURE


@pytest.mark.parametrize("text,secret", [
    ('{"password": "hunter2"}', "hunter2"),
    ("client_secret=abcd1234", "abcd1234"),
    ("access_token=abcd1234&x=1", "abcd1234"),
    ("AWS_SECRET_ACCESS_KEY=abcd1234", "abcd1234"),
    ("X-Api-Key: abcd1234", "abcd1234"),
    ("https://user:pw123@host.test/x", "pw123"),
    ("Set-Cookie: sessionid=abcd1234; Path=/", "abcd1234"),
    ("Cookie: a=abcd1234; b=efgh5678", "efgh5678"),
    ("Authorization: Basic dXNlcjpwYXNzd29yZA==", "dXNlcjpwYXNz"),
    ("Authorization: Bearer abcdefghijkl12345", "abcdefghijkl12345"),
    ("jwt eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiIxMjM0NTY3ODkwIn0.sig-value-here", "eyJzdWIi"),
])
def test_common_credential_formats_are_masked(text, secret):
    assert secret not in san.clean_exception_message(text)


@pytest.mark.parametrize("text", [
    "author: Bob Smith",                       # 'auth' must not match 'author'
    "Message: no such element: Unable to locate element",
    "timeout after 30 seconds waiting for #login",
    "expected 'Welcome back' but got 'Login'",
])
def test_ordinary_diagnostic_text_is_not_altered(text):
    assert san.clean_exception_message(text) == text


@pytest.mark.parametrize("line,masked", [
    ('d.find_element(By.ID,"password").send_keys("hunter2")', True),
    ('d.find_element(By.ID,"user_pass").send_keys("hunter2")', True),
    ('d.find_element(By.ID,"pass").send_keys("hunter2")', True),
    ('d.find_element(By.ID,"passenger").send_keys("Alice")', False),
    ('d.find_element(By.ID,"compass").send_keys("North")', False),
    ('d.find_element(By.ID,"bypass-link").send_keys("x")', False),
])
def test_send_keys_masking_follows_sensitive_names_only(line, masked):
    assert ("***" in san.redact_source_line(line)) is masked


@pytest.mark.parametrize("text,expected", [
    ("C:\\Users\\John Smith\\run_1\\test_a.py:3: in t", "test_a.py:3: in t"),
    ("/home/ann/my project/run/test_b.py:9", "test_b.py:9"),
    ("/tmp/run_9/test_login.py:10: in test_x", "test_login.py:10: in test_x"),
    ('File "/usr/lib/python3/site-packages/selenium/errorhandler.py", line 5', 'File "errorhandler.py", line 5'),
])
def test_absolute_paths_including_spaces_are_reduced_to_basenames(text, expected):
    assert san.strip_paths(text) == expected


def test_legacy_failure_object_in_the_api_is_sanitized_too(client, db_session, test_project, auth_headers):
    run = make_run(db_session, test_project, status="FAILED")
    test = TestResult(run_id=run.id, nodeid="t.py::a", file_name="t.py", method_name="a", status="FAILED")
    db_session.add(test)
    db_session.commit()
    db_session.add(FailureDiagnostic(
        test_result_id=test.id,
        exception_type="selenium.common.exceptions.NoSuchElementException",
        exception_message="Message: x\nStacktrace:\n#0 0x55 <unknown>",
        source_line='d.find_element(By.ID, "password").send_keys("hunter2")',
        raw_traceback="/tmp/run_9/t.py:3: in a\nE   password=hunter2",
    ))
    db_session.commit()

    failure = _results(client, test_project, run, auth_headers).json()["tests"][0]["failure"]

    blob = json.dumps(failure)
    assert "hunter2" not in blob and "/tmp/run_9" not in blob and "0x55" not in blob
    assert failure["source_line"].endswith('send_keys("***")')
    # the stored row itself is unchanged (sanitizing happens on output only)
    stored = db_session.scalar(select(FailureDiagnostic).where(FailureDiagnostic.test_result_id == test.id))
    assert "hunter2" in stored.source_line


def test_report_failure_blocks_are_sanitized_for_every_format(db_session, test_project, tmp_path):
    run = make_run(db_session, test_project, status="RUNNING")
    _persist_failing(db_session, test_project, run, tmp_path)
    failure = db_session.scalar(select(FailureDiagnostic).join(TestResult).where(TestResult.run_id == run.id))
    failure.source_line = 'self.driver.find_element(By.ID, "password").send_keys("hunter2")'
    failure.exception_message = "Message: bad password=hunter2"
    db_session.commit()

    report = build_report_data(db_session, run, test_project)

    assert "hunter2" not in render_html_report(report)
    cells = " ".join(str(c) for row in load_workbook(io.BytesIO(render_excel_report(report)))["Failures"].iter_rows(values_only=True) for c in row)
    assert "hunter2" not in cells
    assert render_pdf_report(report).startswith(b"%PDF")


def test_diagnostic_log_line_is_sanitized():
    from app.services.pytest_diagnostics import TestFailureDiagnostic

    line = TestFailureDiagnostic(
        test_file="t.py", test_class=None, test_method="a", source_line_number=3,
        source_line='d.find_element(By.ID, "password").send_keys("hunter2")',
        exception_type="selenium.common.exceptions.NoSuchElementException",
        exception_message="Message: x password=hunter2\nStacktrace:\n#0 0x55 <unknown>",
    ).format()
    assert "hunter2" not in line and "0x55" not in line and "NoSuchElementException" in line


def test_sanitizer_is_idempotent_and_handles_odd_input():
    raw = "C:\\Users\\bob\\run_1\\test_a.py:3: in t\nE   password: abc"
    once = san.sanitize_traceback(raw)
    assert san.sanitize_traceback(once) == once and "bob" not in once and "abc" not in once
    assert san.sanitize_traceback(None) is None and san.clean_exception_message("") == ""
    assert san.redact_url("not a url at all") == "not a url at all"
    assert san.redact_url("https://user:pw@h.test/p?sig=1&q=2") == "https://***@h.test/p?sig=***&q=2"


# ---------------------------------------------------------------------------
# API + persistence
# ---------------------------------------------------------------------------

LOCATOR_MSG = 'Message: no such element: Unable to locate element: {"method":"css selector","selector":"#missing"}'


def _persist_failing(db_session, project, run, tmp_path, *, with_screenshot=True, password_line=True):
    tmp_path.mkdir(parents=True, exist_ok=True)
    make_project_file(db_session, project, "test_login.py", "# stub")
    make_project_file(db_session, project, "test_ok.py", "# stub")
    source = (
        'self.driver.find_element(By.ID, "password").send_keys("hunter2")' if password_line
        else 'self.driver.find_element(By.ID, "missing").click()'
    )
    trace = (
        "/tmp/run_77/test_login.py:10: in test_x\n"
        f"    {source}\n"
        f"E   selenium.common.exceptions.NoSuchElementException: {LOCATOR_MSG}\n"
        "E     Stacktrace:\nE     #0 0x55 <unknown>"
    )
    junit = tmp_path / "report.xml"
    junit.write_text(
        '<?xml version="1.0"?><testsuites><testsuite tests="2">'
        '<testcase classname="test_login" name="test_x" time="1.0" file="test_login.py">'
        f'<failure message="x">{escape(trace)}</failure></testcase>'
        '<testcase classname="test_ok" name="test_fine" time="0.5" file="test_ok.py"/>'
        "</testsuite></testsuites>", encoding="utf-8",
    )
    shot = "0002_step_2_1700000000000.png"
    steps = [
        {"nodeid": "test_login.py::test_x", "step": 1, "action": "open", "target": "https://app.test/", "value": "",
         "status": "PASS", "duration": 0.1, "error": "", "screenshot": "", "timestamp": 1700000000.0},
        {"nodeid": "test_login.py::test_x", "step": 2, "action": "click", "target": "id=missing", "value": "",
         "status": "FAIL", "duration": 12.0, "error": "no such element", "screenshot": shot if with_screenshot else "",
         "timestamp": 1700000001.0},
    ]
    steps_path = tmp_path / "steps.json"
    steps_path.write_text(json.dumps(steps), encoding="utf-8")
    shots_dir = tmp_path / "screenshots"
    shots_dir.mkdir(exist_ok=True)
    if with_screenshot:
        (shots_dir / shot).write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 16)
    persist_run_results(
        db_session, run, exit_code=1, stdout="", stderr="", junit_report_path=junit,
        steps_json_path=steps_path, screenshots_dir=shots_dir,
    )
    run.status = "FAILED"
    run.completed_at = NOW
    db_session.commit()


def _results(client, project, run, headers):
    return client.get(f"/api/projects/{project.id}/runs/{run.id}/results", headers=headers)


def _failed_and_passed(resp):
    tests = {t["method_name"]: t for t in resp.json()["tests"]}
    return tests["test_x"], tests["test_fine"]


def test_api_returns_analysis_for_failed_and_none_for_passed(client, db_session, test_project, auth_headers, tmp_path):
    run = make_run(db_session, test_project, status="RUNNING", browser="chrome")
    _persist_failing(db_session, test_project, run, tmp_path)

    resp = _results(client, test_project, run, auth_headers)

    assert resp.status_code == 200
    failed, passed = _failed_and_passed(resp)
    a = failed["analysis"]
    assert passed["analysis"] is None
    assert (a["failure_category"], a["confidence"], a["category_basis"]) == ("ELEMENT_NOT_FOUND", "high", "exception_type")
    assert a["exception_type"].endswith("NoSuchElementException")
    assert (a["file_name"], a["line_number"]) == ("test_login.py", 10)
    assert a["failed_step"]["action"] == "click" and a["last_successful_step"]["action"] == "open"
    assert a["page_url"] == "https://app.test/"
    assert a["browser"] == "chrome" and a["timestamp"] is not None
    assert {"label": "Locator", "value": "css selector: #missing"} in a["evidence"]
    assert a["recommended_actions"] and a["likely_cause"]
    assert a["failure_screenshot"]["url"].endswith(f"/screenshots/{a['failure_screenshot']['id']}")
    assert failed["failure"] is not None  # existing contract preserved


def test_api_never_exposes_secrets_or_paths_and_stores_a_sanitized_traceback(client, db_session, test_project, auth_headers, tmp_path):
    run = make_run(db_session, test_project, status="RUNNING")
    _persist_failing(db_session, test_project, run, tmp_path)

    failed, _ = _failed_and_passed(_results(client, test_project, run, auth_headers))

    payload = json.dumps(failed["analysis"]) + str(failed["failure"]["raw_traceback"])
    assert "hunter2" not in payload and "/tmp/run_77" not in payload and "0x55" not in payload
    assert "test_login.py:10: in test_x" in failed["failure"]["raw_traceback"]
    assert failed["analysis"]["traceback"] == failed["failure"]["raw_traceback"]


def test_reading_results_does_not_create_records(client, db_session, test_project, auth_headers, tmp_path):
    run = make_run(db_session, test_project, status="RUNNING")
    _persist_failing(db_session, test_project, run, tmp_path)

    def counts():
        return tuple(db_session.scalar(select(func.count()).select_from(m)) for m in (FailureDiagnostic, TestResult, StepResult, Screenshot))

    before = counts()
    for _ in range(3):
        assert _results(client, test_project, run, auth_headers).status_code == 200
    assert counts() == before


def test_missing_screenshot_gives_empty_state_not_an_error(client, db_session, test_project, auth_headers, tmp_path):
    run = make_run(db_session, test_project, status="RUNNING")
    _persist_failing(db_session, test_project, run, tmp_path, with_screenshot=False)

    failed, _ = _failed_and_passed(_results(client, test_project, run, auth_headers))

    assert failed["analysis"]["failure_screenshot"] is None
    assert {"label": "Failure screenshot", "value": "None recorded"} in failed["analysis"]["evidence"]


def test_legacy_records_with_incomplete_diagnostics(client, db_session, test_project, auth_headers):
    run = make_run(db_session, test_project, status="FAILED")
    sparse = TestResult(run_id=run.id, nodeid="t.py::a", file_name="t.py", method_name="a", status="FAILED")
    empty = TestResult(run_id=run.id, nodeid="t.py::b", file_name="t.py", method_name="b", status="FAILED")
    bare = TestResult(run_id=run.id, nodeid="t.py::c", file_name="t.py", method_name="c", status="FAILED")
    db_session.add_all([sparse, empty, bare])
    db_session.commit()
    db_session.add(FailureDiagnostic(test_result_id=sparse.id, exception_type="selenium.common.exceptions.TimeoutException"))
    db_session.add(FailureDiagnostic(test_result_id=empty.id))
    db_session.commit()

    resp = _results(client, test_project, run, auth_headers)

    assert resp.status_code == 200
    by_name = {t["method_name"]: t["analysis"] for t in resp.json()["tests"]}
    assert by_name["a"]["failure_category"] == "TIMEOUT" and by_name["a"]["traceback"] is None
    assert by_name["b"]["failure_category"] == "UNKNOWN" and by_name["b"]["confidence"] == "low"
    assert by_name["c"]["failure_category"] == "UNKNOWN"


def test_api_ownership_missing_ids_and_auth(client, db_session, test_project, auth_headers, tmp_path):
    run = make_run(db_session, test_project, status="RUNNING")
    _persist_failing(db_session, test_project, run, tmp_path)
    intruder = User(full_name="Intruder", email="intruder.analysis@example.com", mobile_number="9000000044",
                    password_hash=hash_password("x"), is_active=True)
    db_session.add(intruder)
    db_session.commit()
    db_session.refresh(intruder)
    other = {"Authorization": f"Bearer {create_access_token(intruder.id)}"}

    assert _results(client, test_project, run, other).status_code == 404
    assert client.get(f"/api/projects/{test_project.id}/runs/999999999/results", headers=auth_headers).status_code == 404
    assert client.get(f"/api/projects/999999999/runs/{run.id}/results", headers=auth_headers).status_code == 404
    assert client.get(f"/api/projects/{test_project.id}/runs/{run.id}/results").status_code in (401, 403)


def test_running_and_successful_runs_have_no_analysis(client, db_session, test_project, auth_headers, tmp_path):
    queued = make_run(db_session, test_project, status="QUEUED")
    body = _results(client, test_project, queued, auth_headers).json()
    assert body["tests"] is None


# ---------------------------------------------------------------------------
# reports: same analysis in HTML, Excel and PDF
# ---------------------------------------------------------------------------

def _report_data(db_session, project, tmp_path):
    run = make_run(db_session, project, status="RUNNING", browser="chrome")
    _persist_failing(db_session, project, run, tmp_path)
    return run, build_report_data(db_session, run, project)


def test_report_data_carries_the_same_analysis_as_the_api(client, db_session, test_project, auth_headers, tmp_path):
    run, report = _report_data(db_session, test_project, tmp_path)
    api_failed, _ = _failed_and_passed(_results(client, test_project, run, auth_headers))
    report_failed = next(t for t in report.tests if t.method_name == "test_x")
    assert report_failed.analysis.failure_category == api_failed["analysis"]["failure_category"]
    assert report_failed.analysis.summary == api_failed["analysis"]["summary"]
    assert next(t for t in report.tests if t.method_name == "test_fine").analysis is None


def test_html_report_includes_analysis_and_escapes_content(db_session, test_project, tmp_path):
    run, report = _report_data(db_session, test_project, tmp_path)
    failed = next(t for t in report.tests if t.method_name == "test_x")
    object.__setattr__(failed.analysis, "summary", "<script>alert(1)</script>")

    html = render_html_report(report)

    assert "Failure analysis: Element not found" in html and "ELEMENT_NOT_FOUND" in html
    assert "Recommended actions (suggestions)" in html and "hypothesis, not confirmed" in html
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html
    assert "hunter2" not in html and "/tmp/run_77" not in html


def test_excel_report_failures_sheet_has_analysis_columns(db_session, test_project, tmp_path):
    _, report = _report_data(db_session, test_project, tmp_path)
    wb = load_workbook(io.BytesIO(render_excel_report(report)))
    rows = list(wb["Failures"].iter_rows(values_only=True))
    header, first = rows[0], rows[1]
    assert wb.sheetnames == ["Summary", "Tests", "Steps", "Failures", "Screenshots"]
    assert header[:7] == ("Test ID", "Node ID", "Exception Type", "Exception Message", "Source File", "Source Line Number", "Source Line")
    data = dict(zip(header, first))
    assert data["Failure Category"] == "ELEMENT_NOT_FOUND" and data["Confidence"] == "high"
    assert "Locator: css selector: #missing" in data["Evidence"] and data["Recommended Actions"].startswith("1. ")
    assert data["Failed Step"] == "#2 click id=missing" and "hunter2" not in " ".join(str(c) for c in first)


def test_pdf_report_renders_with_analysis(db_session, test_project, tmp_path):
    _, report = _report_data(db_session, test_project, tmp_path)
    pdf = render_pdf_report(report)
    assert pdf.startswith(b"%PDF") and len(pdf) > 1000


def test_report_endpoints_still_work(client, db_session, test_project, auth_headers, tmp_path):
    run, _ = _report_data(db_session, test_project, tmp_path)
    for kind in ("html", "excel", "pdf"):
        resp = client.get(f"/api/projects/{test_project.id}/runs/{run.id}/report/{kind}", headers=auth_headers)
        assert resp.status_code == 200 and len(resp.content) > 500
