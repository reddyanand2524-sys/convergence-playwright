

import base64
import io

import openpyxl
import pytest
from sqlalchemy import event

from app.database import engine
from app.models.failure_diagnostic import FailureDiagnostic
from app.models.run_result import RunResult
from app.models.screenshot import Screenshot
from app.models.step_result import StepResult
from app.models.test_result import TestResult
from app.services.report_data_service import build_report_data
from app.services.report_excel_service import _sanitize_cell_value, render_excel_report
from app.services.report_html_service import render_html_report
from app.services.report_pdf_service import render_pdf_report
from tests.conftest import make_run

_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _make_run_result(db_session, run, **overrides):
    rr = RunResult(
        run_id=run.id,
        exit_code=overrides.get("exit_code", 0),
        total_tests=overrides.get("total_tests", 0),
        passed_count=overrides.get("passed_count", 0),
        failed_count=overrides.get("failed_count", 0),
        skipped_count=overrides.get("skipped_count", 0),
        stdout=overrides.get("stdout", ""),
        stderr=overrides.get("stderr", ""),
    )
    db_session.add(rr)
    db_session.commit()
    db_session.refresh(rr)
    return rr


def _make_test_result(db_session, run, **overrides):
    tr = TestResult(
        run_id=run.id,
        project_file_id=overrides.get("project_file_id"),
        nodeid=overrides.get("nodeid", "test_module.py::test_example"),
        file_name=overrides.get("file_name", "test_module.py"),
        class_name=overrides.get("class_name"),
        method_name=overrides.get("method_name", "test_example"),
        status=overrides.get("status", "PASSED"),
        duration_seconds=overrides.get("duration_seconds", 1.5),
    )
    db_session.add(tr)
    db_session.commit()
    db_session.refresh(tr)
    return tr


def _make_step(db_session, test_result, **overrides):
    step = StepResult(
        test_result_id=test_result.id,
        step_number=overrides.get("step_number", 1),
        action=overrides.get("action", "click"),
        target=overrides.get("target", "#submit"),
        value=overrides.get("value"),
        status=overrides.get("status", "PASS"),
        duration_seconds=overrides.get("duration_seconds", 0.2),
        error=overrides.get("error"),
        sequence_index=overrides.get("sequence_index", 0),
    )
    db_session.add(step)
    db_session.commit()
    db_session.refresh(step)
    return step


def _make_screenshot(db_session, test_result, step_result=None, **overrides):
    shot = Screenshot(
        test_result_id=test_result.id,
        step_result_id=step_result.id if step_result else None,
        file_name=overrides.get("file_name", "shot.png"),
        mime_type=overrides.get("mime_type", "image/png"),
        capture_reason=overrides.get("capture_reason", "ALL_MODE"),
        content=overrides.get("content", _PNG_1PX),
        size_bytes=len(overrides.get("content", _PNG_1PX)),
    )
    db_session.add(shot)
    db_session.commit()
    db_session.refresh(shot)
    return shot


def _make_failure(db_session, test_result, **overrides):
    f = FailureDiagnostic(
        test_result_id=test_result.id,
        exception_type=overrides.get("exception_type", "AssertionError"),
        exception_message=overrides.get("exception_message", "expected true, got false"),
        source_file=overrides.get("source_file", "test_module.py"),
        source_line_number=overrides.get("source_line_number", 42),
        source_line=overrides.get("source_line", "assert result is True"),
        raw_traceback=overrides.get("raw_traceback"),
    )
    db_session.add(f)
    db_session.commit()
    db_session.refresh(f)
    return f


# ---------------------------------------------------------------------------
# 7.1 — Report data service correctness
# ---------------------------------------------------------------------------


def test_report_data_completed_run_basic(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    tr = _make_test_result(db_session, run, status="PASSED", duration_seconds=2.0)

    report = build_report_data(db_session, run, test_project)

    assert report.project.id == test_project.id
    assert report.run.id == run.id
    assert report.summary.total_tests == 1
    assert report.summary.passed_count == 1
    assert len(report.tests) == 1
    assert report.tests[0].id == tr.id
    assert report.tests[0].duration_seconds == 2.0


def test_report_data_failed_run_with_diagnostic(db_session, test_project):
    run = make_run(db_session, test_project, status="FAILED")
    _make_run_result(db_session, run, total_tests=1, failed_count=1)
    tr = _make_test_result(db_session, run, status="FAILED")
    _make_failure(db_session, tr, exception_message="boom")

    report = build_report_data(db_session, run, test_project)

    assert report.tests[0].failure is not None
    assert report.tests[0].failure.exception_message == "boom"
    assert report.summary.failure_count == 1


def test_report_data_zero_tests(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=0)

    report = build_report_data(db_session, run, test_project)

    assert report.tests == []
    assert report.summary.total_tests == 0
    assert report.summary.duration_seconds == 0.0


def test_report_data_missing_run_result_falls_back_to_test_rows(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_test_result(db_session, run, nodeid="a.py::t1", status="PASSED")
    _make_test_result(db_session, run, nodeid="a.py::t2", status="FAILED")

    report = build_report_data(db_session, run, test_project)

    assert report.summary.total_tests == 2
    assert report.summary.passed_count == 1
    assert report.summary.failed_count == 1


def test_report_data_multiple_tests_multiple_steps(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=2, passed_count=2)
    t1 = _make_test_result(db_session, run, nodeid="a.py::t1")
    t2 = _make_test_result(db_session, run, nodeid="a.py::t2")
    _make_step(db_session, t1, sequence_index=0, status="PASS")
    _make_step(db_session, t1, sequence_index=1, status="PASS")
    _make_step(db_session, t2, sequence_index=0, status="FAIL")

    report = build_report_data(db_session, run, test_project)

    by_id = {t.id: t for t in report.tests}
    assert len(by_id[t1.id].steps) == 2
    assert len(by_id[t2.id].steps) == 1
    assert report.summary.total_steps == 3
    assert report.summary.passed_steps == 2
    assert report.summary.failed_steps == 1


def test_report_data_screenshots_step_and_test_level(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1)
    tr = _make_test_result(db_session, run)
    step = _make_step(db_session, tr)
    _make_screenshot(db_session, tr, step_result=step)
    _make_screenshot(db_session, tr, step_result=None)

    report = build_report_data(db_session, run, test_project)

    test_data = report.tests[0]
    assert len(test_data.screenshots) == 1
    assert len(test_data.steps[0].screenshots) == 1
    assert report.summary.screenshot_count == 2


def test_report_data_missing_screenshots_and_diagnostics(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    _make_test_result(db_session, run, status="PASSED")

    report = build_report_data(db_session, run, test_project)

    assert report.tests[0].failure is None
    assert report.tests[0].screenshots == []
    assert report.summary.screenshot_count == 0
    assert report.summary.failure_count == 0


def test_report_data_unicode_and_long_strings(db_session, test_project):
    # nodeid/file_name/method_name are plain VARCHAR columns on SQL Server
    # (not NVARCHAR) — an existing schema characteristic this read-only
    # reporting phase must not change, so unicode here is exercised via the
    # NVARCHAR(max) exception_message column instead, where it round-trips
    # losslessly. Long-string handling is exercised on the same field.
    run = make_run(db_session, test_project, status="FAILED")
    _make_run_result(db_session, run, total_tests=1, failed_count=1)
    long_msg = "x" * 5000 + " 日本語テスト"
    tr = _make_test_result(db_session, run, status="FAILED", nodeid="unicode_test.py::test_long_message")
    _make_failure(db_session, tr, exception_message=long_msg)

    report = build_report_data(db_session, run, test_project)

    assert report.tests[0].failure.exception_message == long_msg
    assert len(report.tests[0].failure.exception_message) == len(long_msg)


def test_report_data_never_fabricates_timestamps_or_traceback(db_session, test_project):
    run = make_run(db_session, test_project, status="FAILED")
    _make_run_result(db_session, run, total_tests=1, failed_count=1)
    tr = _make_test_result(db_session, run, status="FAILED")
    _make_failure(db_session, tr)

    report = build_report_data(db_session, run, test_project)

    assert report.tests[0].started_at is None
    assert report.tests[0].completed_at is None
    assert report.tests[0].failure.raw_traceback is None


# ---------------------------------------------------------------------------
# 7.2 — HTML report
# ---------------------------------------------------------------------------


def test_html_report_contains_expected_content(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    _make_test_result(db_session, run, nodeid="test_module.py::test_example", status="PASSED")

    report = build_report_data(db_session, run, test_project)
    html = render_html_report(report)

    assert f"Run: #{run.id}" in html or f"#{run.id}" in html
    assert "test_module.py::test_example" in html
    assert "PASSED" in html


def test_html_report_escapes_malicious_test_name(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    malicious = "<script>alert(1)</script>"
    _make_test_result(db_session, run, nodeid=malicious, status="PASSED")

    report = build_report_data(db_session, run, test_project)
    html = render_html_report(report)

    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html


def test_html_report_embeds_screenshot_as_base64(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    tr = _make_test_result(db_session, run)
    _make_screenshot(db_session, tr)

    report = build_report_data(db_session, run, test_project)
    html = render_html_report(report)

    expected_b64 = base64.b64encode(_PNG_1PX).decode("ascii")
    assert expected_b64 in html
    assert "data:image/png;base64," in html


def test_html_report_api_endpoint(client, db_session, test_project, auth_headers):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    _make_test_result(db_session, run, status="PASSED")

    resp = client.get(
        f"/api/projects/{test_project.id}/runs/{run.id}/report/html",
        headers=auth_headers,
    )

    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/html")
    assert f"Convergence_Run_{run.id}.html" in resp.headers["content-disposition"]


# ---------------------------------------------------------------------------
# 7.3 — Excel report
# ---------------------------------------------------------------------------


def test_sanitize_cell_value_neutralizes_formula_prefixes():
    assert _sanitize_cell_value("=1+1") == "'=1+1"
    assert _sanitize_cell_value("+cmd") == "'+cmd"
    assert _sanitize_cell_value("-cmd") == "'-cmd"
    assert _sanitize_cell_value("@cmd") == "'@cmd"
    assert _sanitize_cell_value("normal text") == "normal text"
    assert _sanitize_cell_value(42) == 42


def test_excel_report_sheets_and_values(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    tr = _make_test_result(db_session, run, nodeid="test_module.py::test_example", status="PASSED")
    step = _make_step(db_session, tr, action="click", target="#go")
    _make_screenshot(db_session, tr, step_result=step)

    report = build_report_data(db_session, run, test_project)
    content = render_excel_report(report)

    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Summary", "Tests", "Steps", "Failures", "Screenshots"]

    tests_ws = wb["Tests"]
    header = [c.value for c in tests_ws[1]]
    assert header == ["ID", "Node ID", "File", "Class", "Method", "Status", "Duration (s)"]
    row = [c.value for c in tests_ws[2]]
    assert row[1] == "test_module.py::test_example"
    assert row[5] == "PASSED"

    steps_ws = wb["Steps"]
    step_row = [c.value for c in steps_ws[2]]
    assert step_row[3] == "click"
    assert step_row[4] == "#go"

    shots_ws = wb["Screenshots"]
    assert shots_ws.max_row == 2  # header + 1 screenshot


def test_excel_report_formula_injection_neutralized(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    _make_test_result(db_session, run, nodeid="=1+1", status="PASSED")

    report = build_report_data(db_session, run, test_project)
    content = render_excel_report(report)

    wb = openpyxl.load_workbook(io.BytesIO(content))
    tests_ws = wb["Tests"]
    row = [c.value for c in tests_ws[2]]
    assert row[1] == "'=1+1"
    assert not row[1].startswith("=1")


def test_excel_report_api_endpoint(client, db_session, test_project, auth_headers):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    _make_test_result(db_session, run, status="PASSED")

    resp = client.get(
        f"/api/projects/{test_project.id}/runs/{run.id}/report/excel",
        headers=auth_headers,
    )

    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"Convergence_Run_{run.id}.xlsx" in resp.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Summary" in wb.sheetnames


# ---------------------------------------------------------------------------
# 7.4 — PDF report
# ---------------------------------------------------------------------------


def test_pdf_report_magic_bytes_and_nonzero_length(db_session, test_project):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    tr = _make_test_result(db_session, run, status="PASSED")
    _make_screenshot(db_session, tr)

    report = build_report_data(db_session, run, test_project)
    content = render_pdf_report(report)

    assert content.startswith(b"%PDF-")
    assert len(content) > 0


def test_pdf_report_with_failure_diagnostic(db_session, test_project):
    run = make_run(db_session, test_project, status="FAILED")
    _make_run_result(db_session, run, total_tests=1, failed_count=1)
    tr = _make_test_result(db_session, run, status="FAILED")
    _make_failure(db_session, tr, exception_message="<weird> & stuff")

    report = build_report_data(db_session, run, test_project)
    content = render_pdf_report(report)

    assert content.startswith(b"%PDF-")


def test_pdf_report_api_endpoint(client, db_session, test_project, auth_headers):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    _make_test_result(db_session, run, status="PASSED")

    resp = client.get(
        f"/api/projects/{test_project.id}/runs/{run.id}/report/pdf",
        headers=auth_headers,
    )

    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF-")
    assert f"Convergence_Run_{run.id}.pdf" in resp.headers["content-disposition"]


# ---------------------------------------------------------------------------
# 7.5 — auth / ownership / run-state gating
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("report_type", ["html", "excel", "pdf"])
def test_report_requires_auth(client, db_session, test_project, report_type):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=0)

    resp = client.get(f"/api/projects/{test_project.id}/runs/{run.id}/report/{report_type}")
    assert resp.status_code == 401


@pytest.mark.parametrize("report_type", ["html", "excel", "pdf"])
def test_report_404_wrong_project(client, db_session, test_project, auth_headers, report_type):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=0)

    resp = client.get(
        f"/api/projects/999999/runs/{run.id}/report/{report_type}",
        headers=auth_headers,
    )
    assert resp.status_code == 404


@pytest.mark.parametrize("report_type", ["html", "excel", "pdf"])
def test_report_404_wrong_run(client, db_session, test_project, auth_headers, report_type):
    resp = client.get(
        f"/api/projects/{test_project.id}/runs/999999/report/{report_type}",
        headers=auth_headers,
    )
    assert resp.status_code == 404


@pytest.mark.parametrize("report_type", ["html", "excel", "pdf"])
def test_report_404_cross_user(
    client, db_session, test_project, report_type, test_user
):
    from app.core.security import create_access_token, hash_password
    from app.models.project import Project
    from app.models.user import User

    import uuid

    # Unique per test run: must not collide with rows already in a shared database.
    unique = uuid.uuid4().hex[:12]
    other_user = User(
        full_name="Other User",
        email=f"phase7-other-user-{unique}@example.com",
        mobile_number="8" + str(uuid.uuid4().int % 10**9).rjust(9, "0"),
        password_hash=hash_password("Password123!"),
        is_active=True,
    )
    db_session.add(other_user)
    db_session.commit()
    db_session.refresh(other_user)
    other_headers = {"Authorization": f"Bearer {create_access_token(other_user.id)}"}

    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=0)

    resp = client.get(
        f"/api/projects/{test_project.id}/runs/{run.id}/report/{report_type}",
        headers=other_headers,
    )
    assert resp.status_code == 404


@pytest.mark.parametrize("run_status", ["QUEUED", "RUNNING", "CANCELLED"])
@pytest.mark.parametrize("report_type", ["html", "excel", "pdf"])
def test_report_409_for_pending_statuses(
    client, db_session, test_project, auth_headers, run_status, report_type
):
    run = make_run(db_session, test_project, status=run_status)

    resp = client.get(
        f"/api/projects/{test_project.id}/runs/{run.id}/report/{report_type}",
        headers=auth_headers,
    )
    assert resp.status_code == 409
    assert run_status in resp.json()["detail"]


@pytest.mark.parametrize("run_status", ["COMPLETED", "FAILED"])
@pytest.mark.parametrize("report_type", ["html", "excel", "pdf"])
def test_report_200_for_reportable_statuses(
    client, db_session, test_project, auth_headers, run_status, report_type
):
    run = make_run(db_session, test_project, status=run_status)
    _make_run_result(db_session, run, total_tests=0)

    resp = client.get(
        f"/api/projects/{test_project.id}/runs/{run.id}/report/{report_type}",
        headers=auth_headers,
    )
    assert resp.status_code == 200


# ---------------------------------------------------------------------------
# No-data-mutation
# ---------------------------------------------------------------------------


def test_report_generation_does_not_mutate_data(client, db_session, test_project, auth_headers):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=1, passed_count=1)
    tr = _make_test_result(db_session, run, status="PASSED")
    step = _make_step(db_session, tr)
    _make_screenshot(db_session, tr, step_result=step)

    before_counts = {
        "runs": db_session.query(RunResult).count(),
        "tests": db_session.query(TestResult).count(),
        "steps": db_session.query(StepResult).count(),
        "screenshots": db_session.query(Screenshot).count(),
    }
    before_test_status = tr.status
    before_run_status = run.status

    for report_type in ("html", "excel", "pdf"):
        resp = client.get(
            f"/api/projects/{test_project.id}/runs/{run.id}/report/{report_type}",
            headers=auth_headers,
        )
        assert resp.status_code == 200

    after_counts = {
        "runs": db_session.query(RunResult).count(),
        "tests": db_session.query(TestResult).count(),
        "steps": db_session.query(StepResult).count(),
        "screenshots": db_session.query(Screenshot).count(),
    }
    db_session.refresh(tr)
    db_session.refresh(run)

    assert before_counts == after_counts
    assert tr.status == before_test_status
    assert run.status == before_run_status


# ---------------------------------------------------------------------------
# Query-count / N+1 instrumentation
# ---------------------------------------------------------------------------


def test_report_query_count_does_not_scale_linearly(client, db_session, test_project, auth_headers):
    run = make_run(db_session, test_project, status="COMPLETED")
    _make_run_result(db_session, run, total_tests=12, passed_count=12)
    for i in range(12):
        tr = _make_test_result(db_session, run, nodeid=f"test_module.py::test_{i}", status="PASSED")
        step = _make_step(db_session, tr, sequence_index=0)
        _make_screenshot(db_session, tr, step_result=step)

    query_count = 0

    def _count(*args, **kwargs):
        nonlocal query_count
        query_count += 1

    event.listen(engine, "before_cursor_execute", _count)
    try:
        resp = client.get(
            f"/api/projects/{test_project.id}/runs/{run.id}/report/html",
            headers=auth_headers,
        )
    finally:
        event.remove(engine, "before_cursor_execute", _count)

    assert resp.status_code == 200
    # Ownership checks + RunResult + TestResult + batched Step/Screenshot/
    # FailureDiagnostic IN-queries — a small constant, not one query per test.
    assert query_count < 15, f"expected a small constant query count, got {query_count}"
