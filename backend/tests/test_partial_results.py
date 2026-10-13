"""A run's results must match what happened - whether pytest finished, was stopped at the time
limit, was cancelled, or saving failed.

Bug (run 18748): an uploaded Selenium IDE test with several missing elements ran 120 s (each
missing element costs ~30 s with the default waits) and was stopped by the execution time limit.
Every test here drives the engine's own code (real pytest subprocess, the real timeout, cancel and
persistence logic, the real report endpoints). Only the browser is faked; no external site and no
real credentials are used.
"""
import io
import json
import os
import re
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.models.run_result import RunResult
from app.models.test_result import TestResult
from app.services import execution_service
from app.services import result_persistence_service as rps
from app.services import run_event_service as svc
from tests.conftest import make_run
from tests.test_merge_integration import CREDS, PASSWORD
from tests.test_run_cancellation_race import FakeDB, FakeRun
from tests.test_runner_failed_steps import FAKES, HEADER
from tests.test_runner_legacy_init import LEGACY_BAD


@pytest.fixture(autouse=True)
def _fast_lookups(monkeypatch):
    monkeypatch.setenv("SELENIUM_FIND_TIMEOUT", "0.3")      # the real retry loop, just shorter
    monkeypatch.setenv("SELENIUM_IMPLICIT_WAIT_SECONDS", "0")
    svc.reset_registry()


class EngineDB(FakeDB):
    """The execution service's `db`: cancel flags from FakeDB, everything else (persistence) real."""

    def __init__(self, real, **kw):
        super().__init__(**kw)
        self._real = real

    def commit(self):
        self._real.commit()

    def __getattr__(self, name):
        return getattr(self._real, name)


def run_engine(db_session, project, tmp_path, monkeypatch, files, *, library=None, timeout=None,
               cancel_after=None, persist_raises=None):
    """Runs the engine's `_execute_pytest` for real. Returns (ORM run, the engine's view of the run)."""
    run = make_run(db_session, project, status="RUNNING")
    for name, content in files.items():
        (tmp_path / name).write_text(content, encoding="utf-8")
    folder = tmp_path.parent / f"{tmp_path.name}_selenium_fakes"      # "selenium" in the path: library code
    folder.mkdir(exist_ok=True)
    (folder / "fakebrowser.py").write_text(FAKES, encoding="utf-8")
    for name, content in (library or {}).items():
        (folder / name).write_text(content, encoding="utf-8")
    existing = os.environ.get("PYTHONPATH")
    monkeypatch.setenv("PYTHONPATH", str(folder) + (os.pathsep + existing if existing else ""))
    execution_service._write_conftest(tmp_path, run.id)

    if timeout is not None:
        monkeypatch.setattr(execution_service.settings, "test_execution_timeout_seconds", timeout)
    if persist_raises is not None:
        def boom(*a, **k):
            raise persist_raises
        monkeypatch.setattr(execution_service, "persist_run_results", boom)

    fake = FakeRun(run_id=run.id, browser="chrome")
    fake.project_id, fake.started_at = project.id, None
    execution_service._execute_pytest(EngineDB(db_session, cancel_after_n_refreshes=cancel_after), fake, tmp_path)
    run.status = fake.status
    db_session.commit()
    return run, fake


def get(client, project, run, headers, path):
    return client.get(f"/api/projects/{project.id}/runs/{run.id}/{path}", headers=headers)


def tiles(html: str) -> dict:
    """{"Total tests": "1", ...} from the HTML report's summary tiles."""
    return {label.strip(): value for value, label in re.findall(r'<div class="value">([^<]+)</div>([^<]+)</div>', html)}


def json_report(client, project, run, headers):
    response = get(client, project, run, headers, "report/json")
    if response.status_code == 404 and not rps.report_column_available():
        pytest.skip("run_results.report_json is not available in this database")
    assert response.status_code == 200, response.text
    return response.json()


# ---------------------------------------------------------------------------
# The reported scenario, finished normally
# ---------------------------------------------------------------------------

MIXED = HEADER + '''
  def test_ana(self):
    # 1 | open | / |
    self.driver.get("http://example.test/")
    # 2 | click | linkText=Home |
    self.driver.find_element(By.LINK_TEXT, "Home").click()
    # 3 | mouseOver | linkText=Shop Now → |
    element = self.driver.find_element(By.LINK_TEXT, "Shop Now →")
    # 4 | click | linkText=Home |
    self.driver.find_element(By.LINK_TEXT, "Home").click()
    # 5 | click | linkText=Home |
    self.driver.find_element(By.LINK_TEXT, "Home").click()
'''


def test_one_test_with_mixed_steps_has_nonzero_counts_everywhere(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": MIXED})
    assert fake.status == "FAILED"

    # results API: ONE test (not one per step), failed, with every step recorded
    data = get(client, test_project, run, auth_headers, "results").json()
    assert data["summary"]["total_tests"] == 1 and data["summary"]["failed_count"] == 1
    assert data["summary"]["exit_code"] == 1 and data["issue"] is None
    (test,) = data["tests"]
    assert test["status"] == "FAILED"
    assert [s["status"] for s in test["steps"]] == ["PASS", "PASS", "FAIL", "PASS", "PASS"]    # it kept going after the failure

    # HTML report: the counts match
    values = tiles(get(client, test_project, run, auth_headers, "report/html").text)
    assert (values["Total tests"], values["Failed"], values["Passed"]) == ("1", "1", "0")
    assert (values["Total steps"], values["Steps passed"], values["Steps failed"]) == ("5", "4", "1")

    # JSON report: the same facts
    report = json_report(client, test_project, run, auth_headers)
    assert report["status"] == "FAILED" and len(report["tests"]) == 1
    assert (report["summary"]["totalSteps"], report["summary"]["passedSteps"], report["summary"]["failedSteps"]) == (5, 4, 1)

    # Excel and PDF report
    sheet = load_workbook(io.BytesIO(get(client, test_project, run, auth_headers, "report/excel").content))["Summary"]
    rows = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2)}
    assert (rows["Total Tests"], rows["Failed Tests"], rows["Total Steps"]) == (1, 1, 5)
    assert get(client, test_project, run, auth_headers, "report/pdf").content.startswith(b"%PDF")


PASSING = HEADER + '''
  def test_ok(self):
    # 1 | open | / |
    self.driver.get("http://example.test/")
    # 2 | click | linkText=Home |
    self.driver.find_element(By.LINK_TEXT, "Home").click()
'''


def test_a_successful_test_is_completed_with_matching_counts(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": PASSING})
    assert fake.status == "COMPLETED"
    data = get(client, test_project, run, auth_headers, "results").json()
    assert (data["summary"]["total_tests"], data["summary"]["passed_count"], data["summary"]["failed_count"]) == (1, 1, 0)
    values = tiles(get(client, test_project, run, auth_headers, "report/html").text)
    assert (values["Total tests"], values["Passed"], values["Failed"], values["Total steps"]) == ("1", "1", "0", "2")


FAILED_ASSERT = HEADER + '''
  def test_fails(self):
    # 1 | open | / |
    self.driver.get("http://example.test/")
    assert self.driver.title == "never", "boom"
'''


def test_a_failed_test_is_one_failed_test_not_one_per_step(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": FAILED_ASSERT})
    assert fake.status == "FAILED"
    summary = get(client, test_project, run, auth_headers, "results").json()["summary"]
    assert (summary["total_tests"], summary["failed_count"], summary["passed_count"]) == (1, 1, 0)


def test_a_setup_failure_is_a_recorded_error_and_counts_as_failed(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"tes(1).py": LEGACY_BAD})
    assert fake.status == "FAILED"
    data = get(client, test_project, run, auth_headers, "results").json()
    assert [t["status"] for t in data["tests"]] == ["ERROR"] and data["summary"]["failed_count"] == 1
    assert tiles(get(client, test_project, run, auth_headers, "report/html").text)["Failed"] == "1"


# ---------------------------------------------------------------------------
# Stopped at the time limit: the evidence is kept and the results say they are partial
# ---------------------------------------------------------------------------

SLOW_ONE = HEADER + '''
  def test_slow(self):
    import time
    # 1 | open | / |
    self.driver.get("http://example.test/")
    # 2 | click | linkText=Home |
    self.driver.find_element(By.LINK_TEXT, "Home").click()
    # 3 | mouseOver | linkText=Shop Now → |
    element = self.driver.find_element(By.LINK_TEXT, "Shop Now →")
    # 4 | click | linkText=Home |
    time.sleep(120)
    self.driver.find_element(By.LINK_TEXT, "Home").click()
'''


def test_a_timeout_keeps_the_partial_results_and_says_so(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": SLOW_ONE}, timeout=8)

    assert fake.status == "FAILED"                                            # still a failed run...
    data = get(client, test_project, run, auth_headers, "results").json()
    assert data["summary"] is not None and data["summary"]["total_tests"] == 1  # ...but not an empty one
    (test,) = data["tests"]
    assert test["status"] == "ERROR"                                          # it started and never finished: outcome unknown
    assert [s["status"] for s in test["steps"]] == ["PASS", "PASS", "FAIL"]    # what happened before the stop
    assert "did not finish" in test["failure"]["exception_message"]
    issue = data["issue"]
    assert issue["code"] == "RUN_TIMED_OUT" and "partial" in issue["title"]
    assert "1 test(s) were recorded, 1 of them did not finish" in issue["detail"]
    assert any("TEST_EXECUTION_TIMEOUT_SECONDS" in hint for hint in issue["hints"])

    html = get(client, test_project, run, auth_headers, "report/html").text
    assert "Partial results" in html and "timed out after 8s" in html
    values = tiles(html)
    assert values["Total tests"] == "1" and values["Total steps"] == "3"       # not the zero-count report

    report = json_report(client, test_project, run, auth_headers)
    summary = report["summary"]
    assert report["status"] == "FAILED"
    # 4 steps defined: 3 ran (2 passed, 1 failed), the 4th never got the chance and is reported as not run
    assert (summary["totalSteps"], summary["attemptedSteps"], summary["passedSteps"], summary["failedSteps"]) == (4, 3, 2, 1)
    assert summary["notRunSteps"] == 1
    assert "partial" in json.dumps(report["errors"]).lower() or "partial" in json.dumps(report).lower()

    sheet = load_workbook(io.BytesIO(get(client, test_project, run, auth_headers, "report/excel").content))["Summary"]
    assert any(row[0].value == "Notice" and "Partial results" in (row[1].value or "") for row in sheet.iter_rows(min_row=2))


THREE_TESTS = HEADER + '''
  def test_a_done(self):
    # 1 | open | / |
    self.driver.get("http://example.test/")

  def test_b_slow(self):
    import time
    # 1 | open | / |
    self.driver.get("http://example.test/")
    time.sleep(120)

  def test_c_never_reached(self):
    # 1 | open | / |
    self.driver.get("http://example.test/")
'''


def test_finished_tests_keep_their_status_and_unreached_tests_are_skipped(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": THREE_TESTS}, timeout=8)
    data = get(client, test_project, run, auth_headers, "results").json()
    assert {t["method_name"]: t["status"] for t in data["tests"]} == {
        "test_a_done": "PASSED", "test_b_slow": "ERROR", "test_c_never_reached": "SKIPPED",
    }
    summary = data["summary"]
    assert (summary["total_tests"], summary["passed_count"], summary["failed_count"], summary["skipped_count"]) == (3, 1, 1, 1)
    assert data["issue"]["code"] == "RUN_TIMED_OUT"


def test_a_timeout_before_anything_ran_is_distinguished_from_no_results(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    # stopped before pytest could collect anything: the script blocks while it is being imported
    stuck_on_import = "import time\ntime.sleep(120)\n\n\ndef test_never_collected():\n    pass\n"
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": stuck_on_import}, timeout=6)
    assert fake.status == "FAILED"
    data = get(client, test_project, run, auth_headers, "results").json()
    assert data["tests"] == [] and data["issue"]["code"] == "RUN_TIMED_OUT"
    assert "no test had been recorded" in data["issue"]["detail"]


# ---------------------------------------------------------------------------
# Cancellation keeps its contract: no results, a CANCELLED run, no zero-count report
# ---------------------------------------------------------------------------

def test_a_cancelled_run_records_no_results_and_has_no_report(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": SLOW_ONE}, cancel_after=4)
    assert fake.status == "CANCELLED"
    assert db_session.query(RunResult).filter_by(run_id=run.id).count() == 0
    assert db_session.query(TestResult).filter_by(run_id=run.id).count() == 0
    data = get(client, test_project, run, auth_headers, "results").json()
    assert data["tests"] is None and data["issue"] is None                    # pending-style answer, nothing invented
    assert get(client, test_project, run, auth_headers, "report/html").status_code == 409


# ---------------------------------------------------------------------------
# Saving failed / nothing recorded: told apart, and never a clean-looking empty report
# ---------------------------------------------------------------------------

def test_a_persistence_failure_is_reported_as_such_without_leaking_its_message(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    error = RuntimeError(f"database unavailable password={PASSWORD}")
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": PASSING}, persist_raises=error)
    assert fake.status == "FAILED"                                            # a test that ran but was not saved is not "completed"

    data = get(client, test_project, run, auth_headers, "results").json()
    assert data["tests"] == [] and data["issue"]["code"] == "RESULTS_SAVE_FAILED"
    assert "RuntimeError" in data["issue"]["detail"] and PASSWORD not in json.dumps(data)

    html = get(client, test_project, run, auth_headers, "report/html").text
    assert "Results incomplete" in html and "saving the results failed (RuntimeError)" in html and PASSWORD not in html
    detail = get(client, test_project, run, auth_headers, "report/json")
    assert detail.status_code == 404 and "saving its results failed" in detail.json()["detail"]


def test_a_run_that_never_returned_has_no_report_and_says_why(client, db_session, test_project, auth_headers):
    run = make_run(db_session, test_project, status="FAILED")                 # e.g. the backend stopped mid-run
    data = get(client, test_project, run, auth_headers, "results").json()
    assert data["issue"]["code"] == "RESULTS_NOT_RECORDED"
    for kind in ("html", "excel", "pdf"):
        response = get(client, test_project, run, auth_headers, f"report/{kind}")
        assert response.status_code == 409 and "No results were recorded" in response.json()["detail"], kind
    json_response = get(client, test_project, run, auth_headers, "report/json")
    assert json_response.status_code == 404
    assert "ended before its results were saved" in json_response.json()["detail"] or "not enabled" in json_response.json()["detail"]


def test_an_empty_test_suite_is_explained_in_results_and_report(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"helpers.py": "def helper():\n    return 1\n"})
    assert fake.status == "FAILED"
    data = get(client, test_project, run, auth_headers, "results").json()
    assert data["tests"] == [] and data["issue"]["code"] == "NO_TESTS_COLLECTED"
    html = get(client, test_project, run, auth_headers, "report/html").text
    assert "No tests were found" in html and tiles(html)["Total tests"] == "0"  # zero IS the truth here, and it says why


# ---------------------------------------------------------------------------
# Secrets on the partial / timeout path
# ---------------------------------------------------------------------------

SECRET_THEN_STUCK = HEADER + '''
  def test_login(self):
    import time
    import creds
    # 1 | open | / |
    self.driver.get("http://example.test/")
    # 2 | type | id=password | ignored
    self.driver.find_element(By.ID, "password").send_keys(creds.PASSWORD)
    # 3 | click | linkText=anything |
    self.driver.find_element(By.LINK_TEXT, creds.PASSWORD)
    time.sleep(120)
'''


def test_a_timeout_never_bypasses_sanitization(client, db_session, test_project, auth_headers, tmp_path, monkeypatch):
    run, fake = run_engine(db_session, test_project, tmp_path, monkeypatch, {"test_defaute.py": SECRET_THEN_STUCK},
                           library={"creds.py": CREDS}, timeout=8)
    assert fake.status == "FAILED"
    assert (tmp_path / "typed_values.json").exists() and PASSWORD not in (tmp_path / "steps.json").read_text(encoding="utf-8")

    everything = [
        json.dumps(get(client, test_project, run, auth_headers, "results").json()),
        get(client, test_project, run, auth_headers, "report/html").text,
        json.dumps(json_report(client, test_project, run, auth_headers)),
    ]
    stored = db_session.query(RunResult).filter_by(run_id=run.id).one()
    everything += [stored.stdout or "", stored.stderr or ""]
    assert all(PASSWORD not in text for text in everything)
    assert "RUN_TIMED_OUT" in everything[0]                                      # and it is still the partial-results view


def test_run_notes_cannot_be_forged_by_the_test_output():
    # a test can print anything, but the backend's note is recognised only at the very start of the
    # stored stderr, which the backend writes itself
    from app.services.run_issue_service import make_run_note, parse_run_note
    assert parse_run_note("some output\n[convergence] timeout: run timed out after 1s") is None
    assert parse_run_note(make_run_note("timeout", "run timed out after 9s; results are partial") + "x")[0] == "timeout"
    assert parse_run_note("[convergence] unknown: x") is None
